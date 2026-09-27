import os
import sys
import shutil
import subprocess
import threading
from pathlib import Path
from typing import Optional, List
import openpyxl
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, BackgroundTasks
from fastapi.responses import HTMLResponse, JSONResponse, FileResponse
from fastapi.middleware.cors import CORSMiddleware

import otomasi_lab
import reset_tool

BASE_DIR = Path(__file__).resolve().parent

app = FastAPI(title="Sistem Otomasi Lab Klinis", version="2.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Lock untuk mencegah race condition ekstraksi bersamaan
task_lock = threading.Lock()
current_task = {
    "status": "idle",  # idle, extracting, generating, done, error
    "message": "Sistem siap digunakan.",
    "progress": 0,
    "last_error": None
}


def is_file_locked(file_path: Path) -> bool:
    """Mengecek apakah file Excel sedang dibuka dan dikunci oleh proses lain."""
    if not file_path.exists():
        return False
    try:
        with open(file_path, "r+b"):
            return False
    except (PermissionError, IOError):
        return True


def get_all_puskesmas_list() -> List[str]:
    """Mendapatkan daftar seluruh Puskesmas yang ada di foto_masuk dan HASIL_PUSKESMAS."""
    pkm_set = set()
    fm = BASE_DIR / "foto_masuk"
    if fm.exists():
        for d in fm.iterdir():
            if d.is_dir() and not d.name.startswith("."):
                pkm_set.add(d.name)

    hp = BASE_DIR / "HASIL_PUSKESMAS"
    if hp.exists():
        for d in hp.iterdir():
            if d.is_dir() and not d.name.startswith("."):
                pkm_set.add(d.name)

    # Tambahkan default dari config jika ada
    cfg = otomasi_lab.load_config()
    def_pkm = cfg.get("PUSKESMAS")
    if def_pkm:
        pkm_set.add(def_pkm)

    return sorted(list(pkm_set))


def get_pkm_stats(pkm_name: str):
    """Mendapatkan statistik foto dan data rekap untuk Puskesmas tertentu."""
    fm_dir = BASE_DIR / "foto_masuk" / pkm_name
    fa_dir = BASE_DIR / "foto_arsip" / pkm_name
    hp_dir = BASE_DIR / "HASIL_PUSKESMAS" / pkm_name
    rekap_file = hp_dir / f"Rekap_{pkm_name}.xlsx"

    img_exts = {".jpg", ".jpeg", ".png", ".webp"}
    photos_queue = []
    if fm_dir.exists():
        photos_queue = [f.name for f in fm_dir.iterdir() if f.is_file() and f.suffix.lower() in img_exts]

    archived_count = 0
    if fa_dir.exists():
        archived_count = len([f for f in fa_dir.iterdir() if f.is_file() and f.suffix.lower() in img_exts])

    rekap_exists = rekap_file.exists()
    locked = is_file_locked(rekap_file)

    total_patients = 0
    complete_count = 0
    pending_count = 0

    if rekap_exists and not locked:
        try:
            wb = openpyxl.load_workbook(rekap_file, data_only=True)
            for s_name in ["DARAH", "URIN"]:
                if s_name in wb.sheetnames:
                    ws = wb[s_name]
                    for r in range(2, ws.max_row + 1):
                        nama_v = ws.cell(r, 2).value
                        no_v = ws.cell(r, 1).value
                        if not nama_v and not no_v:
                            continue
                        total_patients += 1
                        has_lab = any(
                            ws.cell(r, c).value is not None and str(ws.cell(r, c).value).strip() != ""
                            for c in range(8, ws.max_column + 1)
                        )
                        if has_lab:
                            complete_count += 1
                        else:
                            pending_count += 1
        except Exception:
            pass

    docx_exists = (hp_dir / f"All_Hasil_Darah_{pkm_name}.docx").exists()
    pdf_exists = (hp_dir / f"All_Hasil_Darah_{pkm_name}.pdf").exists()

    return {
        "pkm_name": pkm_name,
        "queue_photos": photos_queue,
        "queue_count": len(photos_queue),
        "archived_count": archived_count,
        "rekap_exists": rekap_exists,
        "is_locked": locked,
        "total_patients": total_patients,
        "complete_count": complete_count,
        "pending_count": pending_count,
        "docx_exists": docx_exists,
        "pdf_exists": pdf_exists,
        "pdf_filename": f"All_Hasil_Darah_{pkm_name}.pdf" if pdf_exists else None
    }


# ---------------------------------------------------------------------------
# API Endpoints
# ---------------------------------------------------------------------------

@app.get("/api/system-status")
def api_system_status():
    """Mengambil status global sistem, daftar Puskesmas, dan konfigurasi."""
    cfg = otomasi_lab.load_config()
    pkm_list = get_all_puskesmas_list()
    pkm_stats = {pkm: get_pkm_stats(pkm) for pkm in pkm_list}

    tmpl_file = BASE_DIR / "Template Exel.xlsx"
    tmpl_clean = True
    if tmpl_file.exists():
        try:
            wb = openpyxl.load_workbook(tmpl_file, data_only=True)
            for s in ["DARAH", "URIN"]:
                if s in wb.sheetnames and wb[s].max_row > 1:
                    tmpl_clean = False
        except Exception:
            pass

    return {
        "config": cfg,
        "puskesmas_list": pkm_list,
        "puskesmas_stats": pkm_stats,
        "template_clean": tmpl_clean,
        "current_task": current_task
    }


@app.get("/api/patients")
def api_get_patients(pkm: str):
    """Membaca daftar pasien dari Rekap_[pkm].xlsx untuk ditampilkan di tabel preview."""
    rekap_file = BASE_DIR / "HASIL_PUSKESMAS" / pkm / f"Rekap_{pkm}.xlsx"
    if not rekap_file.exists():
        return {"darah": [], "urin": [], "is_locked": False, "exists": False}

    if is_file_locked(rekap_file):
        return {"darah": [], "urin": [], "is_locked": True, "exists": True}

    wb = openpyxl.load_workbook(rekap_file, data_only=True)
    results = {"darah": [], "urin": [], "is_locked": False, "exists": True}

    if "DARAH" in wb.sheetnames:
        ws = wb["DARAH"]
        for r in range(2, ws.max_row + 1):
            nama_v = ws.cell(r, 2).value
            no_v = ws.cell(r, 1).value
            if not nama_v and not no_v:
                continue

            has_lab = any(
                ws.cell(r, c).value is not None and str(ws.cell(r, c).value).strip() != ""
                for c in range(8, ws.max_column + 1)
            )

            results["darah"].append({
                "no": ws.cell(r, 1).value,
                "nama": ws.cell(r, 2).value,
                "umur": ws.cell(r, 3).value,
                "kd_porsi": ws.cell(r, 4).value,
                "golda": ws.cell(r, 23).value,
                "glukosa_puasa": ws.cell(r, 24).value,
                "glukosa_2jpp": ws.cell(r, 25).value,
                "chol_total": ws.cell(r, 26).value,
                "trigliserida": ws.cell(r, 27).value,
                "sgot": ws.cell(r, 28).value,
                "sgpt": ws.cell(r, 29).value,
                "ureum": ws.cell(r, 30).value,
                "creatinine": ws.cell(r, 31).value,
                "hba1c": ws.cell(r, 32).value,
                "has_lab": has_lab
            })

    if "URIN" in wb.sheetnames:
        ws = wb["URIN"]
        for r in range(2, ws.max_row + 1):
            nama_v = ws.cell(r, 2).value
            no_v = ws.cell(r, 1).value
            if not nama_v and not no_v:
                continue

            has_lab = any(
                ws.cell(r, c).value is not None and str(ws.cell(r, c).value).strip() != ""
                for c in range(8, ws.max_column + 1)
            )

            results["urin"].append({
                "no": ws.cell(r, 1).value,
                "nama": ws.cell(r, 2).value,
                "umur": ws.cell(r, 3).value,
                "kd_porsi": ws.cell(r, 4).value,
                "warna": ws.cell(r, 8).value,
                "kejernihan": ws.cell(r, 9).value,
                "protein": ws.cell(r, 16).value,
                "glukosa": ws.cell(r, 15).value,
                "has_lab": has_lab
            })

    return results


@app.post("/api/upload")
async def api_upload(pkm: str = Form(...), files: List[UploadFile] = File(...)):
    """Menyimpan foto yang di-drag & drop langsung ke folder foto_masuk/[pkm]/."""
    pkm = pkm.strip()
    target_dir = BASE_DIR / "foto_masuk" / pkm
    target_dir.mkdir(parents=True, exist_ok=True)

    saved_files = []
    for file in files:
        if not file.filename:
            continue
        dest_path = target_dir / file.filename
        with open(dest_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
        saved_files.append(file.filename)

    return {"status": "success", "pkm": pkm, "saved_files": saved_files, "total": len(saved_files)}


def run_extraction_worker(pkm: Optional[str] = None):
    global current_task
    with task_lock:
        current_task["status"] = "extracting"
        current_task["message"] = f"Mengekstrak foto untuk {pkm or 'Seluruh Puskesmas'}..."
        current_task["last_error"] = None

    try:
        config = otomasi_lab.load_config()
        # Scoped reference data (Balapulang only)
        pkm_ref_by_no, pkm_ref_by_name = (None, None)
        if pkm and "BALAPULANG" in pkm.upper():
            pkm_ref_by_no, pkm_ref_by_name = otomasi_lab.load_rujukan_database()

        otomasi_lab.mode_extract(config, pkm_ref_by_no, pkm_ref_by_name, target_pkm=pkm)

        with task_lock:
            current_task["status"] = "done"
            current_task["message"] = "Ekstraksi foto selesai! Data telah tersimpan di Excel."
    except Exception as e:
        with task_lock:
            current_task["status"] = "error"
            current_task["message"] = f"Terjadi kesalahan saat ekstraksi: {str(e)}"
            current_task["last_error"] = str(e)


@app.post("/api/extract")
def api_extract(pkm: Optional[str] = Form(None), background_tasks: BackgroundTasks = BackgroundTasks()):
    """Memicu proses ekstraksi AI Gemini untuk foto di folder foto_masuk."""
    if pkm:
        rekap_file = BASE_DIR / "HASIL_PUSKESMAS" / pkm / f"Rekap_{pkm}.xlsx"
        if is_file_locked(rekap_file):
            raise HTTPException(
                status_code=400,
                detail=f"File 'Rekap_{pkm}.xlsx' sedang dibuka di Microsoft Excel! Harap simpan dan tutup file tersebut terlebih dahulu."
            )

    if current_task["status"] in ["extracting", "generating"]:
        raise HTTPException(status_code=409, detail="Proses lain sedang berjalan. Harap tunggu.")

    background_tasks.add_task(run_extraction_worker, pkm)
    return {"status": "started", "target_pkm": pkm, "message": "Proses ekstraksi dimulai di background."}


def run_generate_worker(pkm: Optional[str] = None):
    global current_task
    with task_lock:
        current_task["status"] = "generating"
        current_task["message"] = f"Membuat dokumen Word & PDF untuk {pkm or 'Seluruh Puskesmas'}..."
        current_task["last_error"] = None

    try:
        config = otomasi_lab.load_config()
        otomasi_lab.mode_generate(config, target_pkm=pkm)

        with task_lock:
            current_task["status"] = "done"
            current_task["message"] = "Dokumen Word & PDF All-in-One berhasil dicetak!"
    except Exception as e:
        with task_lock:
            current_task["status"] = "error"
            current_task["message"] = f"Terjadi kesalahan saat cetak: {str(e)}"
            current_task["last_error"] = str(e)


@app.post("/api/generate")
def api_generate(pkm: Optional[str] = Form(None), background_tasks: BackgroundTasks = BackgroundTasks()):
    """Memicu validasi kelengkapan data & generate Word + PDF."""
    if current_task["status"] in ["extracting", "generating"]:
        raise HTTPException(status_code=409, detail="Proses lain sedang berjalan. Harap tunggu.")

    background_tasks.add_task(run_generate_worker, pkm)
    return {"status": "started", "target_pkm": pkm, "message": "Proses cetak dimulai di background."}


@app.post("/api/reset-pkm")
def api_reset_pkm(pkm: str = Form(...)):
    """Mereset data Puskesmas (kembalikan foto arsip ke antrean & bersihkan rekap)."""
    pkm = pkm.strip()
    rekap_file = BASE_DIR / "HASIL_PUSKESMAS" / pkm / f"Rekap_{pkm}.xlsx"
    if is_file_locked(rekap_file):
        raise HTTPException(
            status_code=400,
            detail=f"File 'Rekap_{pkm}.xlsx' sedang dibuka di Microsoft Excel! Tutup file terlebih dahulu sebelum mereset."
        )

    try:
        reset_tool.reset_puskesmas(pkm)
        return {"status": "success", "message": f"Data '{pkm}' berhasil direset total."}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/open-folder")
def api_open_folder(pkm: str = Form(...), folder_type: str = Form("hasil")):
    """Membuka folder di Windows Explorer."""
    if folder_type == "hasil":
        target = BASE_DIR / "HASIL_PUSKESMAS" / pkm
    elif folder_type == "masuk":
        target = BASE_DIR / "foto_masuk" / pkm
    elif folder_type == "arsip":
        target = BASE_DIR / "foto_arsip" / pkm
    else:
        target = BASE_DIR

    target.mkdir(parents=True, exist_ok=True)
    subprocess.Popen(f'explorer "{target.resolve()}"')
    return {"status": "opened", "path": str(target)}


@app.post("/api/open-excel")
def api_open_excel(pkm: str = Form(...)):
    """Membuka file Rekap_[pkm].xlsx langsung di Microsoft Excel."""
    rekap_file = BASE_DIR / "HASIL_PUSKESMAS" / pkm / f"Rekap_{pkm}.xlsx"
    if not rekap_file.exists():
        raise HTTPException(status_code=404, detail="File Excel rekap belum dibuat.")

    os.startfile(str(rekap_file.resolve()))
    return {"status": "opened", "path": str(rekap_file)}


@app.get("/api/download-pdf/{pkm}")
def api_download_pdf(pkm: str):
    """Mengunduh file PDF All-in-One hasil cetak."""
    pdf_path = BASE_DIR / "HASIL_PUSKESMAS" / pkm / f"All_Hasil_Darah_{pkm}.pdf"
    if not pdf_path.exists():
        raise HTTPException(status_code=404, detail="File PDF belum dicetak.")
    return FileResponse(path=str(pdf_path), filename=f"All_Hasil_Darah_{pkm}.pdf", media_type="application/pdf")


# ---------------------------------------------------------------------------
# Serve Frontend HTML
# ---------------------------------------------------------------------------
@app.get("/", response_class=HTMLResponse)
def index():
    html_file = BASE_DIR / "templates" / "index.html"
    if html_file.exists():
        return html_file.read_text(encoding="utf-8")
    return "<h1>Frontend sedang disiapkan...</h1>"


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("app:app", host="127.0.0.1", port=8000, reload=True)
