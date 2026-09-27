import sys
import shutil
from pathlib import Path
import openpyxl

BASE_DIR = Path(__file__).resolve().parent

def clean_template():
    tmpl = BASE_DIR / "Template Exel.xlsx"
    if tmpl.exists():
        wb = openpyxl.load_workbook(tmpl)
        for s in ["DARAH", "URIN"]:
            if s in wb.sheetnames:
                ws = wb[s]
                if ws.max_row >= 2:
                    ws.delete_rows(2, ws.max_row - 1)
        wb.save(tmpl)
        print(" [OK] 'Template Exel.xlsx' berhasil disterilkan (100% header only).")
    else:
        print(" [WARN] 'Template Exel.xlsx' tidak ditemukan.")

def reset_puskesmas(nama_pkm):
    nama_pkm = nama_pkm.strip()
    if not nama_pkm:
        print(" [ERROR] Nama Puskesmas tidak boleh kosong.")
        return

    print(f"\n >>> MERESET DATA: {nama_pkm} <<<")
    
    # 1. Hapus Rekap Excel di HASIL_PUSKESMAS/[nama_pkm]/
    hasil_dir = BASE_DIR / "HASIL_PUSKESMAS" / nama_pkm
    if hasil_dir.exists():
        for f in hasil_dir.glob("*.xlsx"):
            try:
                f.unlink()
                print(f" [OK] Dihapus: {f.name}")
            except Exception as e:
                print(f" [WARN] Gagal menghapus {f.name}: {e}")
        # Hapus docx dan pdf jika ada
        for f in list(hasil_dir.glob("*.docx")) + list(hasil_dir.glob("*.pdf")):
            try:
                f.unlink()
                print(f" [OK] Dihapus dokumen: {f.name}")
            except Exception as e:
                pass
        # Hapus subfolder FOLDER_PASIEN di dalam hasil jika ada
        hasil_fp = hasil_dir / "FOLDER_PASIEN"
        if hasil_fp.exists():
            shutil.rmtree(str(hasil_fp), ignore_errors=True)

    # 2. Hapus Folder Pasien di FOLDER_PASIEN/[nama_pkm]/
    fp_dir = BASE_DIR / "FOLDER_PASIEN" / nama_pkm
    if fp_dir.exists():
        shutil.rmtree(str(fp_dir), ignore_errors=True)
        print(f" [OK] Folder pasien di 'FOLDER_PASIEN/{nama_pkm}/' telah dibersihkan.")

    # 3. Kembalikan foto dari foto_arsip/[nama_pkm]/ ke foto_masuk/[nama_pkm]/
    arsip_dir = BASE_DIR / "foto_arsip" / nama_pkm
    masuk_dir = BASE_DIR / "foto_masuk" / nama_pkm
    masuk_dir.mkdir(parents=True, exist_ok=True)

    moved = 0
    if arsip_dir.exists():
        for img in list(arsip_dir.glob("*.*")):
            if img.is_file():
                dest = masuk_dir / img.name
                shutil.move(str(img), str(dest))
                moved += 1
        print(f" [OK] {moved} file foto dikembalikan dari arsip ke 'foto_masuk/{nama_pkm}/'")

    # 4. Pastikan Template Exel.xlsx tetap bersih
    clean_template()
    print(f"\n [SELESAI] Data '{nama_pkm}' telah direset total. Meja kerja siap digunakan kembali!\n")

def health_check():
    print("\n" + "=" * 55)
    print("           RINGKASAN STATUS SISTEM LAB")
    print("=" * 55)
    
    # Template
    tmpl = BASE_DIR / "Template Exel.xlsx"
    if tmpl.exists():
        try:
            wb = openpyxl.load_workbook(tmpl, data_only=True)
            r_darah = wb["DARAH"].max_row if "DARAH" in wb.sheetnames else 0
            r_urin = wb["URIN"].max_row if "URIN" in wb.sheetnames else 0
            status_tmpl = "BERSIH (Header Only)" if r_darah <= 1 and r_urin <= 1 else f"KOTOR ({r_darah-1} baris)"
            print(f" [TEMPLATE] Template Exel.xlsx : {status_tmpl}")
        except Exception:
            print(" [TEMPLATE] Template Exel.xlsx : Sedang dibuka di Excel")
    else:
        print(" [TEMPLATE] Template Exel.xlsx : Tidak ditemukan")

    # Foto Masuk
    f_masuk = BASE_DIR / "foto_masuk"
    print("\n [FOTO MASUK (Standby)]:")
    if f_masuk.exists():
        subfolders = [d for d in f_masuk.iterdir() if d.is_dir()]
        for sf in subfolders:
            n_photos = len([f for f in sf.iterdir() if f.is_file()])
            print(f"  -> {sf.name}: {n_photos} foto")
    
    # Foto Arsip
    f_arsip = BASE_DIR / "foto_arsip"
    print("\n [FOTO ARSIP (Sudah Diproses)]:")
    if f_arsip.exists():
        subfolders = [d for d in f_arsip.iterdir() if d.is_dir()]
        for sf in subfolders:
            n_photos = len([f for f in sf.iterdir() if f.is_file()])
            print(f"  -> {sf.name}: {n_photos} foto")

    # Hasil Rekap
    f_hasil = BASE_DIR / "HASIL_PUSKESMAS"
    print("\n [HASIL REKAP EXCEL]:")
    if f_hasil.exists():
        subfolders = [d for d in f_hasil.iterdir() if d.is_dir()]
        for sf in subfolders:
            rekap = sf / f"Rekap_{sf.name}.xlsx"
            if rekap.exists():
                try:
                    wb = openpyxl.load_workbook(rekap, data_only=True)
                    r_darah = max(0, wb["DARAH"].max_row - 1) if "DARAH" in wb.sheetnames else 0
                    r_urin = max(0, wb["URIN"].max_row - 1) if "URIN" in wb.sheetnames else 0
                    print(f"  -> {sf.name}: {r_darah} pasien Darah, {r_urin} pasien Urin")
                except Exception:
                    print(f"  -> {sf.name}: File sedang dibuka di Excel")
            else:
                print(f"  -> {sf.name}: Belum ada file Rekap")
    print("=" * 55 + "\n")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        cmd = sys.argv[1].lower()
        if cmd == "--clean-template":
            clean_template()
        elif cmd == "--reset-pkm" and len(sys.argv) > 2:
            reset_puskesmas(sys.argv[2])
        elif cmd == "--health":
            health_check()
    else:
        health_check()
